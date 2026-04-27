"""
MLA 9th Edition formatter.

Layout:
  - No separate title page.
  - First page starts with a left-aligned block (student name / instructor / course / date),
    then a centered plain-text title, then the essay body immediately follows.
  - Page numbers: "LastName #" in top-right header on every page.
  - Works Cited page: "Works Cited" centered (not bold), hanging indent, double-spaced.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt

from models import EssayJSON
from .base import (
    add_body,
    apply_document_defaults,
    add_page_numbers,
    add_references_page,
    doc_to_bytes,
    extract_last_name,
    set_run_font,
    setup_section,
)


def build(essay: EssayJSON, config: dict) -> bytes:
    doc = Document()
    apply_document_defaults(doc, config)
    section = doc.sections[0]
    setup_section(section, config)

    last_name = extract_last_name(essay.author_name_placeholder)
    add_page_numbers(section, config, last_name=last_name)  # "LastName #" on every page
    _first_page_block(doc, essay, config)
    add_body(doc, essay.body_markdown, config, _heading)
    doc.add_page_break()
    add_references_page(doc, essay.references, config)

    return doc_to_bytes(doc)


# ── First-page header block ──────────────────────────────────────────────────

def _first_page_block(doc: Document, essay: EssayJSON, config: dict) -> None:
    # Top-left four-line block: name, instructor, course, date
    for text in [
        essay.author_name_placeholder,
        essay.instructor_placeholder,
        essay.course_placeholder,
        essay.date,
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = Inches(0)
        run = p.add_run(text)
        set_run_font(run, config)

    # Title: centered, plain text — no bold, no italic, no quotes
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    title_para.paragraph_format.space_before = Pt(0)
    title_para.paragraph_format.space_after = Pt(0)
    title_para.paragraph_format.first_line_indent = Inches(0)
    run = title_para.add_run(essay.title)
    set_run_font(run, config)


# ── Headings ─────────────────────────────────────────────────────────────────

def _heading(doc: Document, text: str, level: int, config: dict) -> None:
    """MLA does not mandate a specific heading style — left-aligned, plain text."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    set_run_font(run, config)
