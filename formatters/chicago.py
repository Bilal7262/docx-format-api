"""
Chicago Author-Date formatter.

Layout:
  - Title page: title centered ~1/3 down the page, then author / course / date.
    NO page number on the title page (different_first_page_header_footer=True).
  - Body: double-spaced, 0.5" first-line indent, centered bold headings.
  - References page: "References" centered (not bold), single-spaced entries
    with a blank line between each, hanging indent.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt

from models import EssayJSON
from .base import (
    add_body,
    apply_document_defaults,
    add_page_numbers,
    add_references_page,
    centered_line,
    doc_to_bytes,
    set_run_font,
    setup_section,
)


def build(essay: EssayJSON, config: dict) -> bytes:
    doc = Document()
    apply_document_defaults(doc, config)
    section = doc.sections[0]
    setup_section(section, config)

    # Suppress page number on the title page; default header shows it on all other pages
    section.different_first_page_header_footer = True
    add_page_numbers(section, config)

    _title_page(doc, essay, config)
    doc.add_page_break()
    add_body(doc, essay.body_markdown, config, _heading)
    doc.add_page_break()
    add_references_page(doc, essay.references, config)

    return doc_to_bytes(doc)


# ── Title page ───────────────────────────────────────────────────────────────

def _title_page(doc: Document, essay: EssayJSON, config: dict) -> None:
    # Title centered roughly one-third down the page, plain text (not bold)
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_before = Inches(3)
    title_para.paragraph_format.space_after = Pt(0)
    title_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    title_para.paragraph_format.first_line_indent = Inches(0)
    run = title_para.add_run(essay.title)
    set_run_font(run, config)

    for text in [
        essay.author_name_placeholder,
        essay.course_placeholder,
        essay.date,
    ]:
        centered_line(doc, text, config)


# ── Headings ─────────────────────────────────────────────────────────────────

def _heading(doc: Document, text: str, level: int, config: dict) -> None:
    """
    Chicago does not rigidly define heading styles for student papers.
    Centered bold for level 1, left-aligned bold for level 2 is widely accepted.
    """
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0)

    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = p.add_run(text)
    run.bold = True
    set_run_font(run, config)
